"""Unit tests for DOCX format preservation v2."""

from io import BytesIO

import pytest

from app.services.latex.docx_builder import build_docx_bytes
from app.services.latex.docx_field_index import build_docx_field_index
from app.services.latex.docx_preserver import preserve_docx_format


def _make_simple_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Senior Engineer at Acme Corp")
    doc.add_paragraph("Built scalable APIs with Python")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_table_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Jane Doe")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Python"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def resume_info():
    return {
        "name": "Jane Doe",
        "headline": "Senior Engineer at Acme Corp",
        "past_experience": [
            {
                "company_name": "Acme Corp",
                "role": "Senior Engineer",
                "description": ["Built scalable APIs with Python"],
            }
        ],
    }


@pytest.fixture
def tailored_resume():
    return {
        "name": "Jane Doe",
        "headline": "Staff Engineer — Platform",
        "past_experience": [
            {
                "company_name": "Acme Corp",
                "role": "Staff Engineer",
                "description": ["Led **FastAPI** platform APIs"],
            }
        ],
    }


def test_build_docx_field_index_simple():
    try:
        import docx  # noqa: F401
    except ImportError:
        pytest.skip("python-docx not installed")

    docx_bytes = _make_simple_docx()
    info = {
        "name": "Jane Doe",
        "headline": "Senior Engineer at Acme Corp",
        "past_experience": [{"description": ["Built scalable APIs with Python"]}],
    }
    index = build_docx_field_index(docx_bytes, info)
    assert index is not None
    assert index["has_tables"] is False
    assert len(index["runs"]) >= 2


def test_preserve_simple_docx_in_place(resume_info, tailored_resume):
    try:
        import docx  # noqa: F401
    except ImportError:
        pytest.skip("python-docx not installed")

    docx_bytes = _make_simple_docx()
    index = build_docx_field_index(docx_bytes, resume_info)
    result, status = preserve_docx_format(docx_bytes, resume_info, tailored_resume, index)
    assert status == "preserved"
    assert result is not None

    from docx import Document

    doc = Document(BytesIO(result))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "FastAPI" in full_text or "Staff Engineer" in full_text


def test_table_docx_falls_back(resume_info, tailored_resume):
    try:
        import docx  # noqa: F401
    except ImportError:
        pytest.skip("python-docx not installed")

    docx_bytes = _make_table_docx()
    index = build_docx_field_index(docx_bytes, resume_info)
    result, status = preserve_docx_format(docx_bytes, resume_info, tailored_resume, index)
    assert status == "fallback"
    assert result is None


def test_preserve_docx_invalid_bytes():
    try:
        import docx  # noqa: F401
    except ImportError:
        pytest.skip("python-docx not installed")
    result, status = preserve_docx_format(b"not-a-docx", {"name": "A"}, {"name": "B"})
    assert result is None
    assert status == "fallback"


def test_house_docx_builder_produces_zip():
    resume = {"name": "Test User", "email": "test@example.com"}
    data = build_docx_bytes(resume)
    assert data[:2] == b"PK"
